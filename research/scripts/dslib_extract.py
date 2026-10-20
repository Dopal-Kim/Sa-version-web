"""DSLIB 건설 프로젝트에서 작업(leaf)별 계획 기간 vs 실제 기간을 뽑는다.

입력: data_raw/dslib/DSLIB 3.4/ (Analysis Sheet + Excel/*.xlsx)
출력: sa_ai/data/dslib_activities.csv, sa_ai/data/dslib_projects.csv

- 대상: Sector가 Construction(*)이고 Tracking 이 있는 프로젝트
- 실제 기간: 마지막 추적 시트(TP)에서 Tracking == 'Finished' 인 작업의 Actual Duration
- 기간 단위: 작업일(working day). 'Xd Yh' → X + Y / (그 프로젝트 일 근무시간)
"""
import csv
import re
import sys
import os
from pathlib import Path

import openpyxl

ROOT = Path(os.environ.get("SA_ROOT") or Path(__file__).resolve().parents[2])
DSLIB = ROOT / "data_raw" / "dslib" / "DSLIB 3.4"
OUT = ROOT / "sa_ai" / "data"

DUR_RE = re.compile(r"^\s*(?:(\d+(?:\.\d+)?)d)?\s*(?:(\d+(?:\.\d+)?)h)?\s*$")


def parse_dur(s, hpd):
    """'442d 2h' → 442.25 (hpd=8). 숫자면 그대로(일). 해석 불가면 None."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    m = DUR_RE.match(str(s))
    if not m or (m.group(1) is None and m.group(2) is None):
        return None
    d = float(m.group(1) or 0)
    h = float(m.group(2) or 0)
    return d + h / hpd


def hours_per_day(wb):
    """Agenda 시트의 근무 시간 칸 수(Yes). 없으면 8."""
    if "Agenda" not in wb.sheetnames:
        return 8.0
    n = 0
    for r in wb["Agenda"].iter_rows(min_row=2, values_only=True):
        if r[0] and isinstance(r[0], str) and ":" in r[0] and str(r[1]).strip().lower() == "yes":
            n += 1
    return float(n) if n else 8.0


def header_index(rows, key="ID"):
    for i, r in enumerate(rows[:6]):
        if r and r[0] == key:
            return i
    return None


def leaf_ids(baseline_rows):
    """WBS 가 다른 WBS 의 접두사가 아닌 행 = 실제 작업."""
    hi = header_index(baseline_rows)
    hdr = baseline_rows[hi]
    ci = {h: k for k, h in enumerate(hdr) if h}
    recs = [r for r in baseline_rows[hi + 1:] if r and r[0] is not None]
    wbs = {r[0]: str(r[ci["WBS"]]) for r in recs if r[ci["WBS"]] is not None}
    allw = set(wbs.values())
    leaves = {i for i, w in wbs.items() if not any(o != w and o.startswith(w + ".") for o in allw)}
    info = {r[0]: {"name": r[ci["Name"]], "wbs": wbs.get(r[0]),
                   "pred": r[ci["Predecessors"]] if "Predecessors" in ci else None}
            for r in recs}
    return leaves, info


def last_tp_sheet(wb):
    tps = [n for n in wb.sheetnames if n.startswith("TP") or n.startswith("Project Control")]
    return tps[-1] if tps else None


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    meta = openpyxl.load_workbook(DSLIB / "DSLIB_Analysis_Sheet.xlsx", read_only=True, data_only=True)["DSLIB"]
    rows = [r for r in list(meta.iter_rows(values_only=True))[3:] if r[0]]
    projs = [r for r in rows if str(r[3]).lower().startswith("construction") and r[9] not in (None, "N/A")]
    files = {p.name.split(" ")[0]: p for p in (DSLIB / "Excel").glob("*.xlsx")}

    act_rows, proj_rows = [], []
    for r in projs:
        code, name, sector, kw = r[0], r[1], r[3], r[4]
        f = files.get(code)
        if not f:
            print("파일 없음", code, file=sys.stderr)
            continue
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        hpd = hours_per_day(wb)
        base = list(wb["Baseline Schedule"].iter_rows(values_only=True))
        leaves, info = leaf_ids(base)
        tp = last_tp_sheet(wb)
        if not tp:
            continue
        trows = list(wb[tp].iter_rows(values_only=True))
        hi = header_index(trows)
        if hi is None:
            continue
        ci = {h: k for k, h in enumerate(trows[hi]) if h}
        n_fin = 0
        for t in trows[hi + 1:]:
            if not t or t[0] is None or t[0] not in leaves:
                continue
            planned = parse_dur(t[ci["Duration"]], hpd)
            actual = parse_dur(t[ci["Actual Duration"]], hpd)
            status = t[ci["Tracking"]] if "Tracking" in ci else None
            if status != "Finished" or not planned or planned <= 0 or actual is None:
                continue
            n_fin += 1
            act_rows.append({
                "project": code, "sector": sector, "act_id": t[0],
                "name": info.get(t[0], {}).get("name"), "wbs": info.get(t[0], {}).get("wbs"),
                "planned_d": round(planned, 4), "actual_d": round(actual, 4),
                "ratio": round(actual / planned, 6) if actual > 0 else 0.0,
            })
        proj_rows.append({"project": code, "name": name, "sector": sector, "keywords": kw,
                          "n_leaf": len(leaves), "n_finished": n_fin, "pd_days": r[11],
                          "early_late": r[17], "hpd": hpd, "last_tp": tp, "tracking_auth": r[9]})
        print(f"{code} 작업 {len(leaves)} 완료 {n_fin}", file=sys.stderr)

    with open(OUT / "dslib_activities.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(act_rows[0].keys()))
        w.writeheader()
        w.writerows(act_rows)
    with open(OUT / "dslib_projects.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(proj_rows[0].keys()))
        w.writeheader()
        w.writerows(proj_rows)
    print(f"프로젝트 {len(proj_rows)}, 완료 작업 {len(act_rows)}", file=sys.stderr)


if __name__ == "__main__":
    main()
