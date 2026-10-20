"""단계 B 백테스트: 작업기간 분포 방식별로 프로젝트 공기 예측이 실제와 얼마나 맞는지 비교.

방법(모두 같은 공정 네트워크·같은 표본 수):
  sym15   : 계획 기간 ±15% 삼각분포 (현재 SA 앱 기본값)
  expert  : DSLIB Risk Analysis 시트의 전문가 3점(낙관/최빈/비관 %) 삼각분포
  lopo    : 다른 프로젝트들의 작업 실적/계획 비율을 기간 구간별로 독립 추출 (leave-one-project-out)
  lopo_blk: 다른 프로젝트 하나를 통째로 뽑고 그 프로젝트의 비율에서 추출 (프로젝트 내 상관 유지)

지표: P80 적중(실제 공기 ≤ 예측 P80), P50 편향, CRPS/계획공기
실제 공기 = 계획 PD × (1 + Early/late)  (DSLIB Analysis Sheet)
"""
import csv
import json
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dslib_extract import DSLIB, OUT, hours_per_day, parse_dur, leaf_ids, header_index  # noqa: E402

N = 2000
SEED = 20261020
BUILDING = ("residential", "commercial", "institutional", "industrial")
DUR_BINS = [0, 2, 5, 10, 20, 50, 1e9]

REL_RE = re.compile(r"^\s*(\d+)\s*(FS|SS|FF|SF)\s*([+-].*)?$", re.I)
REL_RE2 = re.compile(r"^\s*(FS|SS|FF|SF)\s*(\d+)\s*([+-].*)?$", re.I)  # 'FS3' 형식(Successors 열)
LAG_RE = re.compile(r"([+-])?\s*(\d+(?:\.\d+)?)\s*([wdh])", re.I)


def parse_lag(s, hpd):
    if not s:
        return 0.0
    sign = -1.0 if s.strip().startswith("-") else 1.0
    tot = 0.0
    for _, v, u in LAG_RE.findall(s):
        v = float(v)
        tot += v * 5 if u.lower() == "w" else v if u.lower() == "d" else v / hpd
    return sign * tot


def parse_rel(tok, hpd):
    m = REL_RE.match(tok) or None
    if m:
        return int(m.group(1)), m.group(2).upper(), parse_lag(m.group(3), hpd)
    m = REL_RE2.match(tok)
    if m:
        return int(m.group(2)), m.group(1).upper(), parse_lag(m.group(3), hpd)
    return None


def cpm(order, preds, dur):
    """dur: {id: ndarray(N)}. preds: {id: [(pid, type, lag)]}. 반환: 프로젝트 공기 ndarray(N)."""
    es, ef = {}, {}
    for j in order:
        d = dur[j]
        s = np.zeros_like(d)
        for i, t, lag in preds.get(j, []):
            if t == "FS":
                c = ef[i] + lag
            elif t == "SS":
                c = es[i] + lag
            elif t == "FF":
                c = ef[i] + lag - d
            else:  # SF
                c = es[i] + lag - d
            s = np.maximum(s, c)
        es[j], ef[j] = s, s + d
    return np.max(np.vstack([ef[j] for j in order]), axis=0)


def topo(nodes, preds):
    indeg = {n: 0 for n in nodes}
    succ = {n: [] for n in nodes}
    for j in nodes:
        for i, _, _ in preds.get(j, []):
            indeg[j] += 1
            succ[i].append(j)
    q = [n for n in nodes if indeg[n] == 0]
    out = []
    while q:
        n = q.pop()
        out.append(n)
        for s in succ[n]:
            indeg[s] -= 1
            if indeg[s] == 0:
                q.append(s)
    return out if len(out) == len(nodes) else None


def load_project(code, f):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    hpd = hours_per_day(wb)
    base = list(wb["Baseline Schedule"].iter_rows(values_only=True))
    leaves, info = leaf_ids(base)
    hi = header_index(base)
    ci = {h: k for k, h in enumerate(base[hi]) if h}
    wbs = {r[0]: str(r[ci["WBS"]]) for r in base[hi + 1:] if r and r[0] is not None and r[ci["WBS"]] is not None}
    kids = {i: [l for l in leaves if wbs.get(l, "").startswith(w + ".")] for i, w in wbs.items() if i not in leaves}
    dur, preds = {}, {}
    for r in base[hi + 1:]:
        if not r or r[0] not in leaves:
            continue
        d = parse_dur(r[ci["Duration"]], hpd)
        dur[r[0]] = d if d is not None else 0.0
        ps = []
        for tok in str(r[ci["Predecessors"]] or "").split(";"):
            rel = parse_rel(tok, hpd) if tok.strip() else None
            if not rel:
                continue
            pid, t, lag = rel
            for p in ([pid] if pid in leaves else kids.get(pid, [])):
                ps.append((p, t, lag))
        preds[r[0]] = ps
    # 전문가 3점(%)
    expert = {}
    if "Risk Analysis" in wb.sheetnames:
        ra = list(wb["Risk Analysis"].iter_rows(values_only=True))
        rh = header_index(ra)
        if rh is not None:
            rc = {h: k for k, h in enumerate(ra[rh]) if h}
            for r in ra[rh + 1:]:
                if r and r[0] in leaves:
                    try:
                        o, m, p = (float(r[rc["Optimistic"]]), float(r[rc["Most Probable"]]), float(r[rc["Pessimistic"]]))
                        if o <= m <= p and p > 0:
                            expert[r[0]] = (o / 100, m / 100, p / 100)
                    except (TypeError, ValueError, KeyError):
                        pass
    return dur, preds, expert


def family(name):
    return re.sub(r"\s*\(\d+\)\s*$", "", str(name)).strip().lower()


def tri(rng, a, m, b, n):
    if b <= a:
        return np.full(n, m)
    return rng.triangular(a, m, b, n)


def bin_of(d):
    return int(np.searchsorted(DUR_BINS, d, side="right") - 1)


def crps(samples, y):
    x = np.sort(samples)
    n = len(x)
    term1 = np.mean(np.abs(x - y))
    term2 = np.sum((2 * np.arange(1, n + 1) - n - 1) * x) / (n * n)
    return term1 - term2


def main():
    rng = np.random.default_rng(SEED)
    acts = list(csv.DictReader(open(OUT / "dslib_activities.csv", encoding="utf-8")))
    projs = {p["project"]: p for p in csv.DictReader(open(OUT / "dslib_projects.csv", encoding="utf-8"))}
    # 정제: 자리표시(계획<1일) · 취소(실제<0.25일인데 계획≥5일) 제외, 건축 섹터만
    clean = [a for a in acts if float(a["planned_d"]) >= 1
             and not (float(a["actual_d"]) < 0.25 and float(a["planned_d"]) >= 5)
             and any(k in a["sector"].lower() for k in BUILDING)]
    by_proj = {}
    for a in clean:
        by_proj.setdefault(a["project"], []).append((bin_of(float(a["planned_d"])), float(a["ratio"])))
    files = {p.name.split(" ")[0]: p for p in (DSLIB / "Excel").glob("*.xlsx")}

    results = []
    samples = {}
    for code, rs in by_proj.items():
        p = projs[code]
        try:
            el = float(p["early_late"])
            pd_days = float(p["pd_days"])
        except ValueError:
            continue
        dur, preds, expert = load_project(code, files[code])
        nodes = list(dur)
        order = topo(nodes, preds)
        if not order:
            print(code, "순환 관계 — 제외", file=sys.stderr)
            continue
        det = cpm(order, preds, {k: np.array([v]) for k, v in dur.items()})[0]
        if not (0.85 <= det / pd_days <= 1.15):  # 네트워크 해석이 원본 계획공기와 안 맞으면 제외
            print(f"{code} CPM {det:.1f} vs PD {pd_days} — 제외", file=sys.stderr)
            continue
        real = pd_days * (1 + el)
        # 학습 풀: 같은 계열(이름에서 '(n)' 번호를 뗀 것이 같은) 프로젝트까지 통째로 제외 — 형제 프로젝트 누설 방지
        others = [c for c in by_proj if family(projs[c]["name"]) != family(p["name"])]
        pool = {b: np.array([r for c in others for bb, r in by_proj[c] if bb == b]) for b in range(len(DUR_BINS) - 1)}
        allpool = np.array([r for c in others for _, r in by_proj[c]])
        out = {"project": code, "sector": p["sector"], "n_act": len(nodes), "pd": pd_days, "real": real, "det": det}
        for meth in ("sym15", "expert", "lopo", "lopo_blk"):
            sd = {}
            if meth == "lopo_blk":
                donors = rng.choice(others, N)
            for k, d in dur.items():
                if d <= 0:
                    sd[k] = np.zeros(N)
                    continue
                if meth == "sym15":
                    sd[k] = tri(rng, d * 0.85, d, d * 1.15, N)
                elif meth == "expert":
                    o, m, pp = expert.get(k, (1.0, 1.0, 1.0))
                    sd[k] = tri(rng, d * o, d * m, d * pp, N)
                elif meth == "lopo":
                    pl = pool[bin_of(d)]
                    pl = pl if len(pl) >= 30 else allpool
                    sd[k] = d * rng.choice(pl, N)
                else:
                    b = bin_of(d)
                    r = np.empty(N)
                    for c in set(donors):
                        idx = donors == c
                        src = [x for bb, x in by_proj[c] if bb == b]
                        src = np.array(src if len(src) >= 3 else [x for _, x in by_proj[c]])
                        r[idx] = rng.choice(src, idx.sum())
                    sd[k] = d * r
            T = cpm(order, preds, sd) * (pd_days / det)  # 원본 PD 척도로 맞춤
            samples[f"{code}|{meth}"] = np.sort(T).astype(np.float32)
            out[f"{meth}_p50"] = float(np.quantile(T, 0.5))
            out[f"{meth}_p80"] = float(np.quantile(T, 0.8))
            out[f"{meth}_hit80"] = int(real <= np.quantile(T, 0.8))
            out[f"{meth}_crps"] = float(crps(T, real) / pd_days)
            out[f"{meth}_pit"] = float(np.mean(T <= real))
        results.append(out)
        print(f"{code} 실제 {real:.0f} / 계획 {pd_days:.0f} | P80 sym {out['sym15_p80']:.0f} exp {out['expert_p80']:.0f} "
              f"lopo {out['lopo_p80']:.0f} blk {out['lopo_blk_p80']:.0f}", file=sys.stderr)

    np.savez_compressed(OUT / "backtest_samples.npz", **samples)
    with open(OUT / "backtest_projects.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(results[0].keys()))
        w.writeheader()
        w.writerows(results)
    summ = {"n_projects": len(results), "n_samples": N, "seed": SEED}
    for meth in ("sym15", "expert", "lopo", "lopo_blk"):
        summ[meth] = {
            "p80_hit_rate": float(np.mean([r[f"{meth}_hit80"] for r in results])),
            "mean_crps_over_pd": float(np.mean([r[f"{meth}_crps"] for r in results])),
            "median_p50_bias": float(np.median([(r[f"{meth}_p50"] - r["real"]) / r["pd"] for r in results])),
            "pit": [r[f"{meth}_pit"] for r in results],
        }
    json.dump(summ, open(OUT / "backtest_summary.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for meth in ("sym15", "expert", "lopo", "lopo_blk"):
        s = summ[meth]
        print(f"{meth:9s} P80적중 {s['p80_hit_rate']:.2f}  CRPS/PD {s['mean_crps_over_pd']:.3f}  P50편향 {s['median_p50_bias']:+.3f}")


if __name__ == "__main__":
    # 자가검증: FS/SS/FF/SF + lag CPM
    _p = {2: [(1, "FS", 1.0)], 3: [(1, "SS", 2.0)], 4: [(2, "FF", 0.0), (3, "SF", 10.0)]}
    _d = {1: np.array([5.0]), 2: np.array([3.0]), 3: np.array([4.0]), 4: np.array([2.0])}
    # 1: 0~5 / 2: FS+1 → 6~9 / 3: SS+2 → 2~6 / 4: FF → 9-2=7, SF → 2+10-2=10 ⇒ 10~12
    assert cpm(topo([1, 2, 3, 4], _p), _p, _d)[0] == 12.0
    assert family("Apartment Building Finishing Works (3)") == family("apartment building finishing works")
    main()
